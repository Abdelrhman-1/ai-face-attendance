from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import attendance_logic
import database

# Light fills matching the app's status color coding (green/orange/blue/red),
# chosen at readable-in-print lightness rather than the app's dark-theme
# accent colors, since spreadsheets are typically viewed on a white background.
STATUS_FILLS = {
    "حاضر": PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid"),
    "متأخر": PatternFill(start_color="FEEBC8", end_color="FEEBC8", fill_type="solid"),
    "معتذر": PatternFill(start_color="BEE3F8", end_color="BEE3F8", fill_type="solid"),
    "غائب": PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid"),
}

HEADER_FILL = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
CENTER = Alignment(horizontal="center", vertical="center")

HEADERS = [
    "الاسم", "رقم الهوية/الوثيقة", "رقم الجوال",
    "الأحد", "وقت حضور الأحد", "الخميس", "وقت حضور الخميس",
]


def _format_time(check_in_time: str | None) -> str:
    """Extract just the HH:MM portion from a SQLite timestamp, or '-' if absent."""
    if not check_in_time:
        return "-"
    return check_in_time.split(" ")[1][:5] if " " in check_in_time else check_in_time


def build_week_workbook(week: dict, status_filter: str = None) -> Workbook:
    """Build a styled .xlsx workbook summarizing one week's attendance.

    Every registered student gets exactly one row, with both lessons' status
    shown side by side, so no student is ever missing from the report --
    matching the requirement that a registered student always appears with
    an explicit status rather than being silently omitted.

    If status_filter is given, only students whose Sunday or Thursday
    status matches it are included, since each row covers two lessons and
    can't be filtered as a single status on its own.
    """
    sunday_roster = {
        r["student_id"]: r for r in attendance_logic.get_lesson_roster(week["sunday_lesson"]["id"])
    }
    thursday_roster = {
        r["student_id"]: r for r in attendance_logic.get_lesson_roster(week["thursday_lesson"]["id"])
    }
    students = database.get_all_students()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"أسبوع {week['start_date']}"
    sheet.sheet_view.rightToLeft = True

    sheet.append(HEADERS)
    for column_index in range(1, len(HEADERS) + 1):
        cell = sheet.cell(row=1, column=column_index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    for student in students:
        sunday = sunday_roster.get(student["id"])
        thursday = thursday_roster.get(student["id"])
        # A lesson that hasn't been closed yet may still have pending
        # students; report their current known status rather than "absent".
        sunday_status = sunday["status"] if sunday else "لم يسجل بعد"
        thursday_status = thursday["status"] if thursday else "لم يسجل بعد"

        if status_filter and status_filter not in (sunday_status, thursday_status):
            continue

        sheet.append([
            student["name"],
            student["national_id"],
            student["phone"] or "-",
            sunday_status,
            _format_time(sunday["check_in_time"] if sunday else None),
            thursday_status,
            _format_time(thursday["check_in_time"] if thursday else None),
        ])

        row_index = sheet.max_row
        sheet.cell(row=row_index, column=4).alignment = CENTER
        sheet.cell(row=row_index, column=6).alignment = CENTER
        if sunday_status in STATUS_FILLS:
            sheet.cell(row=row_index, column=4).fill = STATUS_FILLS[sunday_status]
        if thursday_status in STATUS_FILLS:
            sheet.cell(row=row_index, column=6).fill = STATUS_FILLS[thursday_status]

    for column_index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = max(16, len(header) + 4)

    return workbook


def export_week_to_bytes(week: dict, status_filter: str = None) -> BytesIO:
    """Render a week's workbook to an in-memory buffer, ready to send as a file."""
    workbook = build_week_workbook(week, status_filter=status_filter)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


DAY_LABELS = {"sunday": "الأحد", "thursday": "الخميس"}
METHOD_LABELS = {"face": "الوجه", "manual": "يدوي", "id": "الهوية"}

RANGE_HEADERS = [
    "التاريخ", "اليوم", "الاسم", "رقم الهوية/الوثيقة",
    "رقم الجوال", "الحالة", "وقت الحضور", "الطريقة",
]


def build_range_workbook(start_date: str, end_date: str, status_filter: str = None):
    """Build a long-format report for every lesson within an arbitrary date range.

    Unlike the single-week matrix report (one row per student), a date
    range can span many lessons, so this uses one row per
    student-per-lesson instead, which stays readable regardless of how
    many weeks the range covers.

    If status_filter is given, only rows matching that exact status are
    included -- straightforward here since each row already covers a
    single lesson.
    """
    lessons = database.get_lessons_in_range(start_date, end_date)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{start_date}_الى_{end_date}"[:31]
    sheet.sheet_view.rightToLeft = True

    sheet.append(RANGE_HEADERS)
    for column_index in range(1, len(RANGE_HEADERS) + 1):
        cell = sheet.cell(row=1, column=column_index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    for lesson in lessons:
        for entry in attendance_logic.get_lesson_roster(lesson["id"]):
            if status_filter and entry["status"] != status_filter:
                continue

            sheet.append([
                lesson["lesson_date"],
                DAY_LABELS.get(lesson["day_type"], lesson["day_type"]),
                entry["name"],
                entry["national_id"],
                entry.get("phone") or "-",
                entry["status"],
                _format_time(entry["check_in_time"]),
                METHOD_LABELS.get(entry["method"], "-") if entry["method"] else "-",
            ])
            row_index = sheet.max_row
            if entry["status"] in STATUS_FILLS:
                sheet.cell(row=row_index, column=6).fill = STATUS_FILLS[entry["status"]]

    for column_index, header in enumerate(RANGE_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = max(14, len(header) + 4)

    return workbook


def export_range_to_bytes(start_date: str, end_date: str, status_filter: str = None) -> BytesIO:
    """Render a date-range report to an in-memory buffer, ready to send as a file."""
    workbook = build_range_workbook(start_date, end_date, status_filter=status_filter)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer